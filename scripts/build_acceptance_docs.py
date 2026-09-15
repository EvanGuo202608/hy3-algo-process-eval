#!/usr/bin/env python3
"""Build Word and PDF copies of acceptance Markdown documents."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    ROOT / "ACCEPTANCE_CHECKLIST.md",
    ROOT / "docs/reports/mvp_summary.md",
    ROOT / "docs/reports/final_submission_note.md",
]


def strip_markdown(line: str) -> str:
    line = re.sub(r"`([^`]*)`", r"\1", line)
    line = line.replace("**", "")
    line = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", line)
    return line


def iter_blocks(markdown: str):
    in_code = False
    code_lines = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                yield ("code", "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line:
            yield ("space", "")
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            yield (f"h{min(level, 3)}", strip_markdown(line.lstrip("#").strip()))
        elif line.startswith("|"):
            yield ("para", strip_markdown(line))
        elif line.startswith("- "):
            yield ("bullet", strip_markdown(line[2:]))
        elif re.match(r"^\d+\. ", line):
            yield ("bullet", strip_markdown(re.sub(r"^\d+\. ", "", line)))
        elif line.startswith(">"):
            yield ("quote", strip_markdown(line.lstrip("> ")))
        else:
            yield ("para", strip_markdown(line))


def build_docx(source: Path) -> Path:
    target = source.with_suffix(".docx")
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for kind, text in iter_blocks(source.read_text(encoding="utf-8")):
        if kind == "space":
            continue
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "h3":
            doc.add_heading(text, level=3)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        else:
            doc.add_paragraph(text)
    doc.save(target)
    return target


def build_pdf(source: Path) -> Path:
    target = source.with_suffix(".pdf")
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "CN",
        parent=styles["Normal"],
        fontName="STSong-Light",
        fontSize=10,
        leading=14,
        spaceAfter=4,
    )
    h1 = ParagraphStyle("CNH1", parent=base, fontSize=18, leading=24, spaceBefore=8, spaceAfter=8)
    h2 = ParagraphStyle("CNH2", parent=base, fontSize=14, leading=20, spaceBefore=6, spaceAfter=6)
    h3 = ParagraphStyle("CNH3", parent=base, fontSize=12, leading=18, spaceBefore=4, spaceAfter=4)
    code = ParagraphStyle("Code", parent=base, fontName="Courier", fontSize=8, leading=10)
    story = []
    for kind, text in iter_blocks(source.read_text(encoding="utf-8")):
        escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if kind == "space":
            story.append(Spacer(1, 3))
        elif kind == "h1":
            story.append(Paragraph(escaped, h1))
        elif kind == "h2":
            story.append(Paragraph(escaped, h2))
        elif kind == "h3":
            story.append(Paragraph(escaped, h3))
        elif kind == "bullet":
            story.append(Paragraph("• " + escaped, base))
        elif kind == "code":
            story.append(Paragraph(escaped.replace("\n", "<br/>"), code))
        else:
            story.append(Paragraph(escaped, base))
    doc = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    doc.build(story)
    return target


def main() -> int:
    for source in SOURCES:
        docx = build_docx(source)
        pdf = build_pdf(source)
        print(docx.relative_to(ROOT))
        print(pdf.relative_to(ROOT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

